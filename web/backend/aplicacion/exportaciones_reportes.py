"""Serialización de listas de control; sin acceso HTTP ni persistencia."""

import csv
import html
import io
from datetime import date
from typing import TypedDict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


class FilaListaControl(TypedDict):
    identificacion: str | None
    apellidos: str
    nombres: str
    seccion: str
    ruta: str
    beneficio: str
    estado: str
    columnaServicio: str


def _celda_segura(valor: object) -> str:
    texto = str(valor if valor is not None else "")
    # Excel interpreta estas celdas como fórmulas incluso dentro de un CSV.
    return f"'{texto}" if texto.lstrip().startswith(("=", "+", "-", "@")) else texto


def _descripcion_filtros(filtros: dict[str, str]) -> str:
    """Convierte los parámetros técnicos en una leyenda legible para impresión."""
    etiquetas = {
        "busqueda": "Búsqueda",
        "ruta": "Ruta",
        "seccion": "Sección",
        "estado": "Estado",
        "confirmacion": "Confirmación",
        "asistencia": "Asistencia",
        "beneficio": "Beneficio",
        "asignacion": "Asignación",
    }
    valores = {
        "confirmada": "Confirmó asistencia",
        "sin_confirmar": "Sin confirmación",
        "presente": "Con registro",
        "sin_registro": "Sin registro",
        "beneficiario": "Beneficiario",
        "no_beneficiario": "No beneficiario",
        "con_ruta": "Con ruta asignada",
        "sin_ruta": "Sin ruta asignada",
    }
    partes = []
    for clave, valor in filtros.items():
        if not valor:
            continue
        etiqueta = etiquetas.get(clave, clave)
        texto = valores.get(str(valor), str(valor))
        partes.append(f"{etiqueta}: {texto}")
    return " · ".join(partes) if partes else "Sin filtros adicionales"


def filas_exportacion(filas: list[FilaListaControl]) -> tuple[list[str], list[list[str]]]:
    if not filas:
        return [
            "N°",
            "Identificación",
            "Apellidos",
            "Nombres",
            "Sección",
            "Ruta",
            "Servicio",
            "Estado",
        ], []
    servicio = filas[0]["columnaServicio"]
    columnas = [
        "N°",
        "Identificación",
        "Apellidos",
        "Nombres",
        "Sección",
        "Ruta",
        servicio,
        "Estado",
    ]
    valores = [
        [
            str(indice),
            _celda_segura(fila["identificacion"]),
            _celda_segura(fila["apellidos"]),
            _celda_segura(fila["nombres"]),
            _celda_segura(fila["seccion"]),
            _celda_segura(fila["ruta"]),
            _celda_segura(fila["beneficio"]),
            _celda_segura(fila["estado"]),
        ]
        for indice, fila in enumerate(filas, start=1)
    ]
    return columnas, valores


def csv_lista_control(columnas: list[str], filas: list[list[str]]) -> bytes:
    salida = io.StringIO(newline="")
    escritor = csv.writer(salida)
    escritor.writerow(columnas)
    escritor.writerows(filas)
    return ("\ufeff" + salida.getvalue()).encode("utf-8")


def xlsx_lista_control(
    titulo: str,
    fecha: date,
    columnas: list[str],
    filas: list[list[str]],
    filtros: dict[str, str],
    institucion: dict[str, str],
) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Lista de control"
    hoja.append([institucion["nombre_colegio"]])
    hoja.append([institucion["subtitulo_reportes"]])
    hoja.append([titulo])
    hoja.append([f"Fecha: {fecha.isoformat()} · Total: {len(filas)}"])
    hoja.append([f"Filtros aplicados: {_descripcion_filtros(filtros)}"])
    hoja.append([])
    hoja.append(columnas)
    for fila in filas:
        hoja.append(fila)
    for celda in hoja[1]:
        celda.font = Font(bold=True, size=14)
    for celda in hoja[3]:
        celda.font = Font(bold=True, size=12)
    for celda in hoja[7]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1E4F8A")
    hoja.freeze_panes = "A8"
    hoja.auto_filter.ref = f"A7:H{max(7, len(filas) + 7)}"
    for columna in hoja.columns:
        letra = columna[0].column_letter
        ancho = min(42, max(12, max(len(str(celda.value or "")) for celda in columna) + 2))
        hoja.column_dimensions[letra].width = ancho
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def html_lista_control(
    titulo: str,
    fecha: date,
    columnas: list[str],
    filas: list[list[str]],
    filtros: dict[str, str],
    institucion: dict[str, str],
) -> str:
    encabezados = "".join(f"<th>{html.escape(columna)}</th>" for columna in columnas)
    cuerpo = "".join(
        "<tr>" + "".join(f"<td>{html.escape(valor)}</td>" for valor in fila) + "</tr>"
        for fila in filas
    )
    return f"""<!doctype html><html lang=\"es\"><head><meta charset=\"utf-8\">
    <title>{html.escape(titulo)}</title><style>
    @page {{ size: landscape; margin: 14mm; }} body {{ color:#14213d; font:11px Arial,sans-serif; }}
    .institucion {{ color:#174c85; font-size:14px; font-weight:bold; margin:0 0 3px; }}
    .subtitulo {{ color:#43546b; margin:0 0 12px; }}
    h1 {{ color:#174c85; font-size:18px; margin:0 0 5px; }} p {{ margin:0 0 8px; color:#43546b; }}
    table {{ width:100%; border-collapse:collapse; }} th,td {{ border:1px solid #aab7c6; padding:6px; text-align:left; }}
    th {{ background:#e8f0f8; color:#14213d; }} tr {{ break-inside:avoid; }}
    .firma {{ margin-top:28px; width:240px; border-top:1px solid #4a5568; padding-top:5px; text-align:center; }}
    @media print {{ button {{ display:none; }} }}
    </style></head><body><p class="institucion">{html.escape(institucion["nombre_colegio"])}</p>
    <p class="subtitulo">{html.escape(institucion["subtitulo_reportes"])}</p>
    <h1>{html.escape(titulo)}</h1>
    <p>Servicio: {html.escape(titulo.rsplit(" — ", 1)[-1])} · Fecha: {fecha.isoformat()} · Registros: {len(filas)}</p>
    <p>Filtros aplicados: {html.escape(_descripcion_filtros(filtros))}</p>
    <table><thead><tr>{encabezados}</tr></thead><tbody>{cuerpo}</tbody></table>
    <div class=\"firma\">Responsable de control</div><script>window.addEventListener('load', () => window.print());</script>
    </body></html>"""
